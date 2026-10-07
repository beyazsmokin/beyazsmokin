#!/usr/bin/env python3
"""İhale dokümanlarının yazısını çıkarır; ajanlar (yapay zekâ) bu metni okur.

Kullanım: python3 belge_metni.py <İhaleler/kod/kaynak> [--cikti calisma/00-belgeler.md] [--sinir 400000]

Okur: PDF (pdfplumber ya da pypdf), DOCX, XLSX (openpyxl), TXT, CSV, MD, HTML, XML ve
ZIP içindeki bu dosyalar. Çizimler (DWG, DXF, görseller) burada okunmaz; onları
`proje_oku.py` raporlar. Okunamayan dosya atlanmaz, "okunamadı" diye listelenir (K-4.6).
Toplam metin `--sinir` karakteri aşarsa kesilir ve kesildiği açıkça yazılır.
"""
import argparse
import html
import io
import re
import sys
import zipfile
from pathlib import Path

METIN = {".txt", ".csv", ".md", ".xml", ".json"}
CIZIM = {".dwg", ".dxf", ".jpg", ".jpeg", ".png", ".gif", ".tif", ".tiff", ".bmp"}


def _pdf(veri: bytes) -> tuple[str, int]:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(veri)) as pdf:
            sayfalar = [(p.extract_text() or "") for p in pdf.pages]
        return "\n\n".join(f"[Sayfa {i}]\n{s}" for i, s in enumerate(sayfalar, 1)), len(sayfalar)
    except ImportError:
        pass
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("PDF okumak için `pip install pdfplumber` gerekli")
    r = PdfReader(io.BytesIO(veri))
    return "\n\n".join(f"[Sayfa {i}]\n{p.extract_text() or ''}" for i, p in enumerate(r.pages, 1)), len(r.pages)


def _docx(veri: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(veri)) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<w:tab/>", "\t", xml)
    xml = re.sub(r"</w:tc>", " | ", xml)
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def _xlsx(veri: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[Sayfa: {ws.title}]")
        for satir in ws.iter_rows(values_only=True):
            if any(v not in (None, "") for v in satir):
                out.append(" | ".join("" if v is None else str(v) for v in satir))
    return "\n".join(out)


def _html(veri: bytes) -> str:
    s = veri.decode("utf-8", "replace")
    s = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", "", s)
    s = re.sub(r"(?i)<br\s*/?>|</(p|div|tr|li|h\d)>", "\n", s)
    return html.unescape(re.sub(r"<[^>]+>", " ", s))


def _metin(veri: bytes) -> str:
    for kod in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            return veri.decode(kod)
        except UnicodeDecodeError:
            pass
    return veri.decode("utf-8", "replace")


def _office_cevir(veri: bytes, uz: str) -> bytes | None:
    """Eski .doc/.rtf/.xls dosyasını bilgisayardaki Word/Excel ile DOCX/XLSX'e çevirir (Windows).
    Office yoksa None döner."""
    if sys.platform != "win32":
        return None
    import os
    import subprocess
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        giris, cikis = os.path.join(d, "girdi" + uz), os.path.join(d, "cikti" + (".xlsx" if uz == ".xls" else ".docx"))
        with open(giris, "wb") as f:
            f.write(veri)
        if uz == ".xls":
            ps = ("$x = New-Object -ComObject Excel.Application; $x.DisplayAlerts = $false; "
                  "$b = $x.Workbooks.Open($env:IA_GIRIS, 0, $true); $b.SaveAs($env:IA_CIKIS, 51); $b.Close($false); $x.Quit()")
        else:
            ps = ("$w = New-Object -ComObject Word.Application; $w.DisplayAlerts = 0; "
                  "$d = $w.Documents.Open($env:IA_GIRIS, $false, $true); $d.SaveAs2($env:IA_CIKIS, 16); $d.Close($false); $w.Quit()")
        try:
            subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps], capture_output=True,
                           timeout=180, env={**os.environ, "IA_GIRIS": giris, "IA_CIKIS": cikis},
                           creationflags=0x08000000)
        except (OSError, subprocess.TimeoutExpired):
            return None
        return Path(cikis).read_bytes() if os.path.exists(cikis) else None


def oku(ad: str, veri: bytes) -> tuple[str | None, str]:
    """(metin, not) döner; metin None ise not okunamama nedenidir."""
    uz = Path(ad).suffix.lower()
    try:
        if uz == ".pdf":
            m, n = _pdf(veri)
            bos = not re.sub(r"\[Sayfa \d+\]|\s", "", m)
            return (None, f"{n} sayfa, yazı yok (taranmış PDF; görsel inceleme gerekli)") if bos \
                else (m, f"{n} sayfa")
        if uz == ".docx":
            return _docx(veri), "Word"
        if uz in (".xlsx", ".xlsm"):
            return _xlsx(veri), "Excel"
        if uz in (".html", ".htm"):
            return _html(veri), "HTML"
        if uz in METIN:
            return _metin(veri), "metin"
        if uz in (".doc", ".xls", ".rtf"):
            yeni = _office_cevir(veri, uz)
            if yeni is not None:
                return (_xlsx(yeni), "Excel (eski biçimden çevrildi)") if uz == ".xls"                     else (_docx(yeni), "Word (eski biçimden çevrildi)")
            return None, "eski Office biçimi okunamadı; dosyayı PDF ya da DOCX/XLSX olarak kaydedin"
        if uz in CIZIM:
            return None, "çizim/görsel: metraj ajanı proje_oku.py raporundan inceler"
        return None, "desteklenmeyen biçim"
    except Exception as e:  # bir dosya bozuksa diğerleri yine okunur
        return None, f"okunamadı: {e}"


def _zip_adi(uye: zipfile.ZipInfo) -> str:
    """UTF-8 işareti olmayan zip adları Türkçe DOS kodlamasıyla (cp857) yazılmıştır (EKAP)."""
    if uye.flag_bits & 0x800:
        return uye.filename
    try:
        return uye.filename.encode("cp437").decode("cp857")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return uye.filename


def _zip_uyeleri(z: zipfile.ZipFile, onek: str, derinlik: int = 0):
    """Zip içindeki dosyalar; içteki zip'ler de açılır (EKAP dokümanı zip içinde zip gelir)."""
    for uye in z.infolist():
        if uye.is_dir():
            continue
        ad = f"{onek}/{_zip_adi(uye)}"
        if Path(ad).name.startswith("~$"):  # Word/Excel'in açık dosya kilidi, belge değil
            continue
        veri = z.read(uye)
        if ad.lower().endswith(".zip") and derinlik < 3:
            try:
                with zipfile.ZipFile(io.BytesIO(veri)) as ic:
                    yield from _zip_uyeleri(ic, ad, derinlik + 1)
                continue
            except zipfile.BadZipFile:
                pass
        yield ad, veri


def dosyalar(kaynak: Path):
    """(görünen ad, bayt) üretir; açılmış klasörü olan zip tekrar okunmaz."""
    for f in sorted(p for p in kaynak.rglob("*") if p.is_file() and not p.name.startswith((".", "~$"))):
        goreli = str(f.relative_to(kaynak)).replace("\\", "/")
        if f.suffix.lower() == ".zip":
            if (f.parent / f.stem).is_dir():
                continue
            try:
                with zipfile.ZipFile(f) as z:
                    yield from _zip_uyeleri(z, goreli)
            except zipfile.BadZipFile:
                yield goreli, b""
            continue
        yield goreli, f.read_bytes()


def topla(kaynak: Path, sinir: int = 400_000) -> dict:
    """Bütün dokümanların metnini tek Markdown'da toplar."""
    parcalar, liste, kalan, kesilen, sayfa = [], [], sinir, 0, 0
    for ad, veri in dosyalar(kaynak):
        metin, notu = oku(ad, veri)
        m = re.match(r"(\d+) sayfa", notu)
        sayfa += int(m.group(1)) if m else 0
        liste.append(f"| {ad} | {'okundu' if metin else 'okunamadı'} | {notu} |")
        if not metin:
            continue
        metin = re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", metin)).strip()
        if len(metin) > kalan:
            kesilen += len(metin) - max(kalan, 0)
            metin = metin[:max(kalan, 0)]
        kalan -= len(metin)
        if metin:
            parcalar.append(f"\n## {ad}\n\n{metin}\n")
    md = ["# İhale dokümanları (metin)\n", "| Dosya | Durum | Not |", "|---|---|---|", *liste]
    if kesilen:
        md.append(f"\n> UYARI: toplam metin {sinir} karakter sınırını aştı; {kesilen} karakter "
                  "okunmadı. Eksik kalan bölümler için dosyanın aslına bakın.")
    return {"md": "\n".join(md + parcalar) + "\n", "dosya": len(liste),
            "okunan": sum("| okundu |" in s for s in liste), "sayfa": sayfa, "kesilen": kesilen}


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("kaynak", type=Path)
    p.add_argument("--cikti", type=Path)
    p.add_argument("--sinir", type=int, default=400_000)
    a = p.parse_args()
    s = topla(a.kaynak, a.sinir)
    if a.cikti:
        a.cikti.write_text(s["md"], encoding="utf-8")
        print(f"{s['okunan']}/{s['dosya']} dosya okundu, {s['sayfa']} sayfa -> {a.cikti}")
    else:
        sys.stdout.write(s["md"])


if __name__ == "__main__":
    main()
