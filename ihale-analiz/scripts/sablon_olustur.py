#!/usr/bin/env python3
"""Excel rapor şablonlarını üretir. Şablonun tek kaynağı bu dosyadır.

Kullanım: python3 sablon_olustur.py [cikti-klasoru]   (varsayılan: templates/excel)

Üretir:
  birim-fiyat.xlsx     BFTC'li (birim fiyat teklif cetvelli) ihaleler
  anahtar-teslim.xlsx  Anahtar teslim / götürü bedel ihaleler

Her şablonda gizli "_harita" sayfası, hangi sayfa sütununun hangi analiz
dosyasından doldurulacağını tanımlar; excel_rapor.py yalnızca bu haritaya
göre yazar, şablonun yapısını değiştirmez. Şablonda değişiklik isteniyorsa
bu dosya güncellenip yeniden çalıştırılır.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

SURUM = "1.0"
SATIR = 500                      # veri satırı kapasitesi
BAS = 5                          # ilk veri satırı (1 başlık, 2 açıklama, 3 toplam, 4 sütun adları)
SON = BAS + SATIR - 1

LACIVERT = "163460"
TOPRAK = "C8A57A"
GIRIS = PatternFill("solid", fgColor="FFFFFF")
HESAP = PatternFill("solid", fgColor="EEF1F6")
TOPLAM = PatternFill("solid", fgColor="F3E9DA")
BASLIK = PatternFill("solid", fgColor=LACIVERT)
SUTUN = PatternFill("solid", fgColor="2B4C7E")
INCE = Side(style="thin", color="D0D5DD")
KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
PARA = '#,##0.00'
MIKTAR = '#,##0.000'
YUZDE = '0.0%'

KIRMIZI = PatternFill("solid", fgColor="F4C7C3")
TURUNCU = PatternFill("solid", fgColor="FCD9B6")
SARI = PatternFill("solid", fgColor="FCE8B2")
YESIL = PatternFill("solid", fgColor="D9EAD3")


def baslik(ws, metin, aciklama, genislik):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=genislik)
    c = ws.cell(1, 1, metin)
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = BASLIK
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=genislik)
    c = ws.cell(2, 1, aciklama)
    c.font = Font(italic=True, size=9, color="667085")
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.row_dimensions[2].height = 30


def tablo(ws, sutunlar, toplamlar=()):
    """sutunlar: (başlık, genişlik, biçim, formül şablonu veya None). Formülde {r} satır no."""
    for j, (ad, gen, bicim, formul) in enumerate(sutunlar, 1):
        harf = get_column_letter(j)
        c = ws.cell(4, j, ad)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = SUTUN
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = KENAR
        ws.column_dimensions[harf].width = gen
        for r in range(BAS, SON + 1):
            h = ws.cell(r, j)
            if formul:
                h.value = formul.format(r=r)
                h.fill = HESAP
            if bicim:
                h.number_format = bicim
            h.border = KENAR
            h.alignment = Alignment(vertical="top", wrap_text=bicim is None)
    ws.row_dimensions[4].height = 32
    ws.cell(3, 1, "TOPLAM").font = Font(bold=True)
    for j in range(1, len(sutunlar) + 1):
        ws.cell(3, j).fill = TOPLAM
        ws.cell(3, j).border = KENAR
    for j in toplamlar:
        harf = get_column_letter(j)
        c = ws.cell(3, j, f"=SUM({harf}{BAS}:{harf}{SON})")
        c.font = Font(bold=True)
        c.number_format = sutunlar[j - 1][2] or PARA
    ws.freeze_panes = ws.cell(BAS, 1)
    ws.auto_filter.ref = f"A4:{get_column_letter(len(sutunlar))}{SON}"
    ws.sheet_view.zoomScale = 90
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "4:4"


def secim(ws, sutun, secenekler):
    dv = DataValidation(type="list", formula1='"' + ",".join(secenekler) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{sutun}{BAS}:{sutun}{SON}")


def renk(ws, aralik, kurallar):
    for metin, dolgu in kurallar:
        ws.conditional_formatting.add(
            aralik, FormulaRule(formula=[f'ISNUMBER(SEARCH("{metin}",{aralik.split(":")[0]}))'], fill=dolgu))


TOLERANSLAR = {  # Özet'teki ad -> (tanımlı ad, varsayılan, biçim)
    "Metraj Toleransı": ("MetrajTol", 0.05, YUZDE),
    "Pursantaj ✔ eşiği (puan)": ("PursUyum", 1, "0.00"),
    "Pursantaj ✗ eşiği (puan)": ("PursCiddi", 2.5, "0.00"),
    "Kazı ek derinliği (m)": ("KaziEk", 0.10, "0.00"),
    "Hendek genişliği (m)": ("HendekB", 1.00, "0.00"),
    "İksa derinlik eşiği (m)": ("IksaH", 1.50, "0.00"),
}


def ozet(wb, tur, sonuc_satirlari, toleranslar):
    ws = wb.active
    ws.title = "Özet"
    baslik(ws, f"İHALE ANALİZ RAPORU · {tur}",
           "Beyaz hücreler analizden doldurulur, gri hücreler hesaplanır. Karar ve gerekçe raporun ilk satırıdır.", 5)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 30
    alanlar = [
        ("İhale Kayıt No (İKN)", None), ("İdare", None), ("İşin Adı", None), ("İhale Usulü", None),
        ("Teklif Türü", None), ("İhale Tarihi ve Saati", None), ("Yaklaşık Maliyet", PARA),
        ("Geçici Teminat Oranı", YUZDE), ("Geçici Teminat Tutarı", PARA), ("İşin Süresi", None),
        ("İşin Yeri", None), ("Karar", None), ("Karar Gerekçesi", None),
        *[(ad, TOLERANSLAR[ad][2]) for ad in toleranslar], ("Analiz Tarihi", "dd.mm.yyyy"),
    ]
    r = 4
    ws.cell(3, 1, "İHALE BİLGİLERİ").font = Font(bold=True, color=LACIVERT, size=11)
    satir = {}
    for ad, bicim in alanlar:
        a = ws.cell(r, 1, ad)
        a.font = Font(bold=True)
        a.border = KENAR
        b = ws.cell(r, 2)
        b.border = KENAR
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if bicim:
            b.number_format = bicim
        satir[ad] = r
        r += 1
    ym, oran = satir["Yaklaşık Maliyet"], satir["Geçici Teminat Oranı"]
    ws.cell(satir["Geçici Teminat Tutarı"], 2, f'=IF(OR(B{ym}="",B{oran}=""),"",B{ym}*B{oran})').fill = HESAP
    for ad in toleranslar:
        isim, varsayilan, _ = TOLERANSLAR[ad]
        ws.cell(satir[ad], 2, varsayilan)
        wb.defined_names[isim] = DefinedName(isim, attr_text=f"'Özet'!$B${satir[ad]}")
    k = ws.cell(satir["Karar"], 2)
    k.font = Font(bold=True, size=12)
    dv = DataValidation(type="list", formula1='"Katıl,Şartlı katıl,Katılma"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(k.coordinate)
    ws.conditional_formatting.add(k.coordinate, CellIsRule(operator="equal", formula=['"Katıl"'], fill=YESIL))
    ws.conditional_formatting.add(k.coordinate, CellIsRule(operator="equal", formula=['"Şartlı katıl"'], fill=SARI))
    ws.conditional_formatting.add(k.coordinate, CellIsRule(operator="equal", formula=['"Katılma"'], fill=KIRMIZI))

    wb.defined_names["YaklasikMaliyet"] = DefinedName("YaklasikMaliyet", attr_text=f"'Özet'!$B${ym}")

    ws.cell(3, 4, "SONUÇ ÖZETİ").font = Font(bold=True, color=LACIVERT, size=11)
    r = 4
    for ad, formul, bicim in sonuc_satirlari:
        a = ws.cell(r, 4, ad)
        a.font = Font(bold=True)
        a.border = KENAR
        b = ws.cell(r, 5, formul)
        b.fill = HESAP
        b.border = KENAR
        b.number_format = bicim
        r += 1
    ws.column_dimensions["E"].width = 22
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return satir


def ortak_sayfalar(wb):
    ws = wb.create_sheet("Yeterlilik")
    baslik(ws, "YETERLİLİK KONTROLÜ", "İdari şartnamedeki katılım ve yeterlilik koşullarının firma profiliyle karşılaştırması.", 5)
    tablo(ws, [("Sıra", 6, None, None), ("Koşul", 50, None, None), ("Şartname Maddesi", 16, None, None),
               ("Firma Durumu", 16, None, None), ("Not", 50, None, None)])
    secim(ws, "D", ["Karşılıyor", "Eksik", "Belirsiz"])
    renk(ws, f"D{BAS}:D{SON}", [("Eksik", KIRMIZI), ("Belirsiz", SARI), ("Karşılıyor", YESIL)])

    ws = wb.create_sheet("Teknik")
    baslik(ws, "TEKNİK DEĞERLENDİRME", "Teknik şartname gereksinimleri, karşılanamayan ve rekabeti kısıtlayan maddeler.", 5)
    tablo(ws, [("Sıra", 6, None, None), ("Gereksinim", 50, None, None), ("Şartname Maddesi", 16, None, None),
               ("Durum", 18, None, None), ("Not", 50, None, None)])
    secim(ws, "D", ["Karşılanıyor", "Karşılanamıyor", "Belirsiz", "Kısıtlayıcı"])
    renk(ws, f"D{BAS}:D{SON}", [("Karşılanamıyor", KIRMIZI), ("Kısıtlayıcı", TURUNCU), ("Belirsiz", SARI)])

    ws = wb.create_sheet("Mali")
    baslik(ws, "MALİ DEĞERLENDİRME", "Teminat, ödeme, avans, fiyat farkı, ceza ve kârlılık kalemleri.", 5)
    tablo(ws, [("Sıra", 6, None, None), ("Kalem", 36, None, None), ("Değer", 22, None, None),
               ("Kaynak", 18, None, None), ("Not", 56, None, None)])

    ws = wb.create_sheet("Riskler")
    baslik(ws, "RİSK LİSTESİ", "Risk denetçisinin derecelendirdiği riskler ve önlemler.", 5)
    tablo(ws, [("Sıra", 6, None, None), ("Risk", 50, None, None), ("Kaynak", 18, None, None),
               ("Seviye", 12, None, None), ("Önlem", 50, None, None)])
    secim(ws, "D", ["Düşük", "Orta", "Yüksek"])
    renk(ws, f"D{BAS}:D{SON}", [("Yüksek", KIRMIZI), ("Orta", SARI), ("Düşük", YESIL)])

    ws = wb.create_sheet("Kişisel Hesap")
    baslik(ws, "KİŞİSEL HESAP",
           "Kullanıcının kendi hesap kurallarıyla bulunan tutarın sistem tahminiyle karşılaştırması "
           "(kisisel_hesap.py çıktısı).", 10)
    tablo(ws, [("Sıra", 6, None, None), ("Poz No", 16, None, None), ("Tanım", 40, None, None),
               ("Birim", 8, None, None), ("Miktar", 13, MIKTAR, None), ("Birim Fiyat", 14, PARA, None),
               ("Sistem Tutarı", 17, PARA, None), ("Kişisel Tutar", 17, PARA, None),
               ("Fark", 15, PARA, '=IF(OR(G{r}="",H{r}=""),"",H{r}-G{r})'),
               ("Uygulanan Kurallar", 36, None, None)], toplamlar=(7, 8, 9))

    ws = wb.create_sheet("Yapılacaklar")
    baslik(ws, "TEKLİF ÖNCESİ YAPILACAKLAR", "Teklif verilmeden tamamlanacak işler.", 5)
    tablo(ws, [("Sıra", 6, None, None), ("İş", 56, None, None), ("Son Tarih", 14, "dd.mm.yyyy", None),
               ("Sorumlu", 20, None, None), ("Durum", 14, None, None)])
    secim(ws, "E", ["Bekliyor", "Devam ediyor", "Tamam"])
    renk(ws, f"E{BAS}:E{SON}", [("Tamam", YESIL), ("Devam", SARI)])


ORTAK_HARITA = [
    # sayfa, kaynak, sütun başlığı, alan adları (| ile alternatif)
    ("Özet", "alanlar:01-ozet.md", "İhale Kayıt No (İKN)", "İKN|İhale kayıt no|İhale kodu"),
    ("Özet", "alanlar:01-ozet.md", "İdare", "İdare"),
    ("Özet", "alanlar:01-ozet.md", "İşin Adı", "İşin adı|İşin adı / niteliği / miktarı|İş"),
    ("Özet", "alanlar:01-ozet.md", "İhale Usulü", "İhale usulü"),
    ("Özet", "alanlar:01-ozet.md", "Teklif Türü", "Teklif türü|Teklif türü (birim fiyat / götürü bedel)"),
    ("Özet", "alanlar:01-ozet.md", "İhale Tarihi ve Saati", "İhale tarihi ve saati|İhale tarihi"),
    ("Özet", "alanlar:01-ozet.md", "Yaklaşık Maliyet", "Yaklaşık maliyet|Yaklaşık maliyet (açıklanmışsa)"),
    ("Özet", "alanlar:01-ozet.md", "Geçici Teminat Oranı", "Geçici teminat oranı"),
    ("Özet", "alanlar:01-ozet.md", "İşin Süresi", "İşin süresi|İşin süresi / yeri"),
    ("Özet", "alanlar:01-ozet.md", "İşin Yeri", "İşin yeri"),
    ("Özet", "karar:rapor", "Karar", "karar"),
    ("Özet", "karar:rapor", "Karar Gerekçesi", "gerekce"),
    ("Özet", "karar:rapor", "Analiz Tarihi", "tarih"),
    ("Yeterlilik", "tablo:02-idari.md", "Sıra", "#sira"),
    ("Yeterlilik", "tablo:02-idari.md", "Koşul", "Koşul"),
    ("Yeterlilik", "tablo:02-idari.md", "Şartname Maddesi", "Madde"),
    ("Yeterlilik", "tablo:02-idari.md", "Firma Durumu", "Firma durumu|Firma durumu (Karşılıyor / Eksik / Belirsiz)|Durum"),
    ("Yeterlilik", "tablo:02-idari.md", "Not", "Not"),
    ("Teknik", "tablo:03-teknik.md", "Sıra", "#sira"),
    ("Teknik", "tablo:03-teknik.md", "Gereksinim", "Gereksinim"),
    ("Teknik", "tablo:03-teknik.md", "Şartname Maddesi", "Madde"),
    ("Teknik", "tablo:03-teknik.md", "Durum", "Durum"),
    ("Teknik", "tablo:03-teknik.md", "Not", "Not"),
    ("Mali", "tablo:04-mali.md", "Sıra", "#sira"),
    ("Mali", "tablo:04-mali.md", "Kalem", "Kalem"),
    ("Mali", "tablo:04-mali.md", "Değer", "Değer"),
    ("Mali", "tablo:04-mali.md", "Kaynak", "Kaynak|Madde"),
    ("Mali", "tablo:04-mali.md", "Not", "Not"),
    ("Riskler", "tablo:05-riskler.md", "Sıra", "#sira"),
    ("Riskler", "tablo:05-riskler.md", "Risk", "Risk"),
    ("Riskler", "tablo:05-riskler.md", "Kaynak", "Kaynak"),
    ("Riskler", "tablo:05-riskler.md", "Seviye", "Seviye|Seviye (Düşük/Orta/Yüksek)"),
    ("Riskler", "tablo:05-riskler.md", "Önlem", "Önlem"),
    ("Kişisel Hesap", "csv:kisisel-hesap.csv", "Sıra", "#sira"),
    ("Kişisel Hesap", "csv:kisisel-hesap.csv", "Poz No", "poz_no"),
    ("Kişisel Hesap", "csv:kisisel-hesap.csv", "Tanım", "tanim"),
    ("Kişisel Hesap", "csv:kisisel-hesap.csv", "Birim", "birim"),
    ("Kişisel Hesap", "csv:kisisel-hesap.csv", "Miktar", "miktar"),
    ("Kişisel Hesap", "csv:kisisel-hesap.csv", "Birim Fiyat", "birim_fiyat"),
    ("Kişisel Hesap", "csv:kisisel-hesap.csv", "Sistem Tutarı", "sistem_tutar"),
    ("Kişisel Hesap", "csv:kisisel-hesap.csv", "Kişisel Tutar", "kisisel_tutar"),
    ("Kişisel Hesap", "csv:kisisel-hesap.csv", "Uygulanan Kurallar", "uygulanan_kurallar"),
    ("Yapılacaklar", "csv:yapilacaklar.csv", "Sıra", "#sira"),
    ("Yapılacaklar", "csv:yapilacaklar.csv", "İş", "is"),
    ("Yapılacaklar", "csv:yapilacaklar.csv", "Son Tarih", "son_tarih"),
    ("Yapılacaklar", "csv:yapilacaklar.csv", "Sorumlu", "sorumlu"),
    ("Yapılacaklar", "csv:yapilacaklar.csv", "Durum", "durum"),
]


def harita(wb, satirlar, tur):
    ws = wb.create_sheet("_harita")
    ws.append(["sablon", tur, "surum", SURUM])
    ws.append(["sayfa", "kaynak", "sütun başlığı", "alan"])
    for s in satirlar:
        ws.append(list(s))
    ws.sheet_state = "hidden"


def ek_sayfalar(wb):
    """Fiyat Kaynakları ve Açık Sorular: iki şablonda da bulunur."""
    ws = wb.create_sheet("Fiyat Kaynakları")
    baslik(ws, "FİYAT KAYNAKLARI",
           "Her birim fiyatın hangi listeden, hangi sayfadan ve hangi döneme ait olduğu. Kaynağı yazılmayan fiyat "
           "rapora girmez.", 9)
    tablo(ws, [("Sıra No", 7, None, None), ("Poz No", 16, None, None), ("İmalatın Cinsi", 44, None, None),
               ("Birim", 8, None, None), ("Birim Fiyat", 14, PARA, None), ("Kaynak Listesi", 34, None, None),
               ("Dayanak / Sayfa", 24, None, None), ("Dönem", 22, None, None), ("Not", 30, None, None)])

    ws = wb.create_sheet("Açık Sorular")
    baslik(ws, "AÇIK SORULAR / KARARLAR",
           "Hiçbir kalem eksik kalmasın: kullanıcıya ya da idareye sorulacaklar, tutara etkisi ve verilen cevap.", 6)
    tablo(ws, [("No", 6, None, None), ("Satır / Poz", 26, None, None), ("Soru", 56, None, None),
               ("Etki (TL)", 16, PARA, None), ("Durum", 18, None, None), ("Cevap / Not", 50, None, None)],
          toplamlar=(4,))
    dv = DataValidation(type="list", formula1='"Açık,Karar bekliyor,Cevaplandı"', allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"E{BAS}:E{SON}")
    renk(ws, f"E{BAS}:E{SON}", [("CEVAPLAND", YESIL), ("BEKL", SARI), ("AÇ", TURUNCU)])


EK_HARITA = [
    ("Fiyat Kaynakları", "csv:fiyat-kaynaklari.csv", "Sıra No", "#sira"),
    ("Fiyat Kaynakları", "csv:fiyat-kaynaklari.csv", "Poz No", "poz_no"),
    ("Fiyat Kaynakları", "csv:fiyat-kaynaklari.csv", "İmalatın Cinsi", "tanim"),
    ("Fiyat Kaynakları", "csv:fiyat-kaynaklari.csv", "Birim", "birim"),
    ("Fiyat Kaynakları", "csv:fiyat-kaynaklari.csv", "Birim Fiyat", "birim_fiyat"),
    ("Fiyat Kaynakları", "csv:fiyat-kaynaklari.csv", "Kaynak Listesi", "kaynak"),
    ("Fiyat Kaynakları", "csv:fiyat-kaynaklari.csv", "Dayanak / Sayfa", "dayanak"),
    ("Fiyat Kaynakları", "csv:fiyat-kaynaklari.csv", "Dönem", "donem"),
    ("Fiyat Kaynakları", "csv:fiyat-kaynaklari.csv", "Not", "not"),
    ("Açık Sorular", "csv:acik-sorular.csv", "No", "#sira"),
    ("Açık Sorular", "csv:acik-sorular.csv", "Satır / Poz", "satir"),
    ("Açık Sorular", "csv:acik-sorular.csv", "Soru", "soru"),
    ("Açık Sorular", "csv:acik-sorular.csv", "Etki (TL)", "etki"),
    ("Açık Sorular", "csv:acik-sorular.csv", "Durum", "durum"),
    ("Açık Sorular", "csv:acik-sorular.csv", "Cevap / Not", "cevap"),
]

ACIK_SORU = ("Açık soru (cevaplanmamış)",
             f"=COUNTA('Açık Sorular'!C{BAS}:C{SON})-COUNTIF('Açık Sorular'!E{BAS}:E{SON},\"CEVAPLAND*\")", "0")

SIFIRSIZ = '#,##0.00;-#,##0.00;;@'   # 0 değerini boş gösterir
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım",
         "Aralık"]


def birim_fiyat(cikti: Path):
    wb = Workbook()
    K = "'BFTC Kıyas'"
    ozet(wb, "BİRİM FİYAT (BFTC)", [
        ("İdare Yaklaşık Maliyet", f"={K}!J3", PARA),
        ("Teklif tutarı (hesap miktarı)", f"={K}!K3", PARA),
        ("İdare YM / açıklanan yaklaşık maliyet", f'=IF(OR(YaklasikMaliyet="",{K}!J3=0),"",{K}!J3/YaklasikMaliyet)', YUZDE),
        ("Kalem sayısı", f'=COUNTA({K}!B{BAS}:B{SON})', "0"),
        ("Cetvelde eksik kalem", f'=COUNTIF({K}!L{BAS}:L{SON},"Cetvelde eksik")', "0"),
        ("Cetvelde fazla kalem", f'=COUNTIF({K}!L{BAS}:L{SON},"Cetvelde fazla")', "0"),
        ("Projede bulunamayan", f'=COUNTIF({K}!L{BAS}:L{SON},"Projede bulunamadı")', "0"),
        ("Cetvelde olmayan iş", f'=COUNTIF({K}!L{BAS}:L{SON},"Cetvelde yok")', "0"),
        ("Yüksek risk", f'=COUNTIF(Riskler!D{BAS}:D{SON},"Yüksek")', "0"),
        ("Eksik yeterlilik", f'=COUNTIF(Yeterlilik!D{BAS}:D{SON},"Eksik")', "0"),
        ACIK_SORU,
        ("Sistem tahmini (kişisel hesap)", "='Kişisel Hesap'!G3", PARA),
        ("Sizin yönteminizle", "='Kişisel Hesap'!H3", PARA),
    ], toleranslar=["Metraj Toleransı"])

    ws = wb.create_sheet("BFTC Kıyas")
    baslik(ws, "BİRİM FİYAT TEKLİF CETVELİ · METRAJ KIYASI",
           "İdare miktarı cetvelden, hesaplanan miktar projeden gelir. Durum, Özet sayfasındaki metraj toleransına göre "
           "hesaplanır. Teklif birim fiyatı girildiğinde tutarlar oluşur.", 14)
    tablo(ws, [
        ("Sıra No", 7, None, None),
        ("Poz No", 16, None, None),
        ("İş Kalemi Tanımı", 44, None, None),
        ("Birim", 8, None, None),
        ("İdare Miktarı", 14, MIKTAR, None),
        ("Hesaplanan Miktar", 14, MIKTAR, None),
        ("Fark", 13, MIKTAR, '=IF(OR(E{r}="",F{r}=""),"",F{r}-E{r})'),
        ("Fark %", 9, YUZDE, '=IF(OR(E{r}="",F{r}="",E{r}=0),"",(F{r}-E{r})/E{r})'),
        ("Teklif Birim Fiyatı", 14, PARA, None),
        ("Tutar (İdare Miktarı)", 17, PARA, '=IF(OR(E{r}="",I{r}=""),"",E{r}*I{r})'),
        ("Tutar (Hesap Miktarı)", 17, PARA, '=IF(OR(F{r}="",I{r}=""),"",F{r}*I{r})'),
        ("Durum", 18, None,
         '=IF(B{r}="","",IF(E{r}="","Cetvelde yok",IF(F{r}="","Projede bulunamadı",'
         'IF(H{r}>MetrajTol,"Cetvelde eksik",IF(H{r}<-MetrajTol,"Cetvelde fazla","Uygun")))))'),
        ("Geçmiş Birim Fiyat (Ort.)", 15, PARA, None),
        ("Not", 40, None, None),
    ], toplamlar=(10, 11))
    renk(ws, f"L{BAS}:L{SON}", [("Cetvelde eksik", TURUNCU), ("Cetvelde fazla", SARI),
                                ("Projede bulunamadı", KIRMIZI), ("Cetvelde yok", KIRMIZI), ("Uygun", YESIL)])
    # teklif fiyatı geçmiş ortalamadan %20'den fazla sapıyorsa işaretle
    ws.conditional_formatting.add(
        f"I{BAS}:I{SON}",
        FormulaRule(formula=[f'AND(ISNUMBER(I{BAS}),ISNUMBER(M{BAS}),M{BAS}>0,ABS(I{BAS}/M{BAS}-1)>0.2)'], fill=SARI))

    ek_sayfalar(wb)
    ortak_sayfalar(wb)
    harita(wb, ORTAK_HARITA + EK_HARITA + [
        ("BFTC Kıyas", "csv:metraj-kiyas.csv", "Sıra No", "#sira"),
        ("BFTC Kıyas", "csv:metraj-kiyas.csv", "Poz No", "poz_no"),
        ("BFTC Kıyas", "csv:metraj-kiyas.csv", "İş Kalemi Tanımı", "tanim"),
        ("BFTC Kıyas", "csv:metraj-kiyas.csv", "Birim", "birim"),
        ("BFTC Kıyas", "csv:metraj-kiyas.csv", "İdare Miktarı", "idare"),
        ("BFTC Kıyas", "csv:metraj-kiyas.csv", "Hesaplanan Miktar", "hesap"),
        ("BFTC Kıyas", "csv:metraj-kiyas.csv", "Teklif Birim Fiyatı", "birim_fiyat"),
        ("BFTC Kıyas", "csv:metraj-kiyas.csv", "Geçmiş Birim Fiyat (Ort.)", "gecmis_fiyat"),
        ("BFTC Kıyas", "csv:metraj-kiyas.csv", "Not", "not"),
    ], "birim-fiyat")
    wb.save(cikti / "birim-fiyat.xlsx")


def anahtar_teslim(cikti: Path):
    wb = Workbook()
    M = "'Metraj Mahal Listesi'"
    P = "'Pursantaj Keşif Analizi'"
    ML = lambda c: f"{M}!${c}${BAS}:${c}${SON}"   # noqa: E731
    ozet(wb, "ANAHTAR TESLİM / GÖTÜRÜ BEDEL", [
        ("Metraj toplamı", f"={M}!K3", PARA),
        ("Metraj / yaklaşık maliyet", f'=IF(OR(YaklasikMaliyet="",{M}!K3=0),"",{M}!K3/YaklasikMaliyet)', YUZDE),
        ("BFTC toplamı", "=BFTC!G3", PARA),
        ("Poz sayısı", f"=COUNTA(BFTC!B{BAS}:B{SON})", "0"),
        ("Metraj satırı", f"=COUNTA({ML('C')})", "0"),
        ("Miktarı girilmemiş satır", f'=SUMPRODUCT(({ML("C")}<>"")*(LEFT({ML("C")},1)<>"—")*({ML("F")}=""))', "0"),
        ("Fiyatı olmayan satır", f'=SUMPRODUCT(({ML("F")}<>"")*(LEFT({ML("C")},1)<>"—")*({ML("J")}=0))', "0"),
        ("BEKLİYOR işaretli satır", f'=COUNTIF({ML("I")},"*BEKLİYOR*")', "0"),
        ("Pursantajla eşlenmemiş tutar", f"=IF({M}!K3=0,\"\",{M}!K3-{P}!E3)", PARA),
        ("Pursantaj ✗ ciddi fark", f'=COUNTIF({P}!I{BAS}:I{SON},"✗*")', "0"),
        ("Pursantaj ⚠ dikkat", f'=COUNTIF({P}!I{BAS}:I{SON},"⚠*")', "0"),
        ACIK_SORU,
        ("Yüksek risk", f'=COUNTIF(Riskler!D{BAS}:D{SON},"Yüksek")', "0"),
        ("Eksik yeterlilik", f'=COUNTIF(Yeterlilik!D{BAS}:D{SON},"Eksik")', "0"),
        ("Sistem tahmini (kişisel hesap)", "='Kişisel Hesap'!G3", PARA),
        ("Sizin yönteminizle", "='Kişisel Hesap'!H3", PARA),
    ], toleranslar=["Pursantaj ✔ eşiği (puan)", "Pursantaj ✗ eşiği (puan)", "Kazı ek derinliği (m)",
                    "Hendek genişliği (m)", "İksa derinlik eşiği (m)"])

    ws = wb.create_sheet("BFTC")
    baslik(ws, "BİRİM FİYAT TEKLİF CETVELİ (KENDİ HESABIMIZ)",
           "Miktar elle girilmez: Metraj Mahal Listesi'ndeki aynı pozların toplamıdır (Durum'u KAPSAM DIŞI / DAHİL DEĞİL "
           "ile başlayanlar hariç). "
           "Birim fiyatın kaynağı Fiyat Kaynakları sayfasındadır.", 8)
    tablo(ws, [
        ("Sıra No", 7, None, None),
        ("Poz No", 16, None, None),
        ("İmalatın Cinsi", 50, None, None),
        ("Birim", 8, None, None),
        ("Miktar", 14, MIKTAR,
         f'=IF(B{{r}}="","",SUMIFS({ML("F")},{ML("C")},B{{r}},{ML("I")},"<>KAPSAM DIŞI*",'
         f'{ML("I")},"<>DAHİL DEĞİL*"))'),
        ("Birim Fiyat", 14, PARA, None),
        ("Tutar", 17, PARA, '=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})'),
        ("Not", 34, None, None),
    ], toplamlar=(7,))
    ws.conditional_formatting.add(
        f"F{BAS}:F{SON}", FormulaRule(formula=[f'AND(B{BAS}<>"",F{BAS}="")'], fill=SARI))

    ws = wb.create_sheet("Dönemsel BFTC")
    son_ay = get_column_letter(5 + len(AYLAR) + 2)     # Yıllık Kitap sütunu
    baslik(ws, "DÖNEMSEL BFTC",
           "Aynı pozun aylık ve yıllık liste fiyatları. Kullanılan Dönem seçilir, tutar o dönemin fiyatıyla hesaplanır.",
           5 + len(AYLAR) + 6)
    sut = [("Sıra No", 7, None, None), ("Poz No", 16, None, None), ("İmalatın Cinsi", 40, None, None),
           ("Birim", 8, None, None),
           ("Miktar", 13, MIKTAR, '=IF(B{r}="","",SUMIF(BFTC!$B$5:$B$504,B{r},BFTC!$E$5:$E$504))')]
    sut += [(ay, 11, PARA, None) for ay in AYLAR]
    sut += [("Yıllık Liste", 12, PARA, None), ("Yıllık Kitap", 12, PARA, None),
            ("Kullanılan Dönem", 13, None, None),
            ("Birim Fiyat", 13, PARA,
             f'=IF(OR(B{{r}}="",T{{r}}=""),"",IFERROR(1/(1/INDEX(F{{r}}:{son_ay}{{r}},'
             f'MATCH(T{{r}},$F$4:${son_ay}$4,0))),""))'),
            ("Tutar", 16, PARA, '=IF(OR(E{r}="",U{r}=""),"",E{r}*U{r})'),
            ("Not / Dönem Kaynağı", 30, None, None)]
    tablo(ws, sut, toplamlar=(22,))
    dv = DataValidation(type="list", formula1=f"$F$4:${son_ay}$4", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"T{BAS}:T{SON}")

    ws = wb.create_sheet("Teknik Tarifler")
    baslik(ws, "TEKNİK TARİFLER", "BFTC'deki her pozun teknik tarifi ve ölçü kuralı (birim fiyat kitabından).", 3)
    tablo(ws, [("Sıra No", 7, None, None), ("Poz No", 16, None, None), ("Teknik Tarifi", 120, None, None)])

    ek_sayfalar(wb)       # Fiyat Kaynakları, Açık Sorular; aşağıdaki sayfalar Açık Sorular'ın önüne girer

    ws = wb.create_sheet("Metraj Mahal Listesi", index=wb.sheetnames.index("Açık Sorular"))
    baslik(ws, "METRAJ MAHAL LİSTESİ",
           "Her satır bir pozun bir iş grubundaki miktarıdır; kaynağı (çizim, katman, etiket) ve hesap kuralı yazılır. "
           "BFTC miktarları ve pursantaj payları buradan hesaplanır. Durum'da 'BEKLİYOR' geçen satırlar eksik sayılır; "
           "Durum'u 'KAPSAM DIŞI' ya da 'DAHİL DEĞİL' ile başlayan satırlar tutara ve BFTC miktarına girmez. Poz No'su "
           "'—' ile başlayan satırlar bilgi/not satırıdır, eksik sayılmaz.", 11)
    tablo(ws, [
        ("İş Grubu", 18, None, None),
        ("Alt Başlık", 22, None, None),
        ("Poz No", 15, None, None),
        ("İmalatın Cinsi", 40, None, None),
        ("Birim", 7, None, None),
        ("Miktar", 13, MIKTAR, None),
        ("Kaynak (çizim / katman / etiket)", 34, None, None),
        ("Hesap (ölçü kuralı)", 34, None, None),
        ("Durum", 24, None, None),
        ("Birim Fiyat (BFTC'den)", 14, SIFIRSIZ,
         '=IF(C{r}="",0,IFERROR(VLOOKUP(C{r},BFTC!$B$5:$F$504,5,FALSE),0))'),
        ("Satır Tutarı", 16, SIFIRSIZ,
         '=IF(OR(F{r}="",J{r}=0,LEFT(I{r},11)="KAPSAM DIŞI",LEFT(I{r},11)="DAHİL DEĞİL"),0,'
         'F{r}*J{r})'),
    ], toplamlar=(11,))
    renk(ws, f"I{BAS}:I{SON}", [("BEKLİYOR", SARI), ("KAPSAM DIŞI", KIRMIZI), ("DAHİL DEĞİL", KIRMIZI),
                                ("HESAPLANDI", YESIL), ("DOĞRULANDI", YESIL)])

    ws = wb.create_sheet("Mevcut Mahal Listesi", index=wb.sheetnames.index("Açık Sorular"))
    baslik(ws, "MEVCUT MAHAL LİSTESİ",
           "Sahada mevcut olan ve yeni imalata dahil olmayan tesisler. Kapsam dışı bırakmanın şartname dayanağı "
           "Durum sütununa yazılır.", 8)
    tablo(ws, [("Sıra No", 7, None, None), ("Grup", 22, None, None), ("İmalat / Mevcut Tesis", 40, None, None),
               ("Birim", 8, None, None), ("Miktar", 13, MIKTAR, None), ("Ölçüm Kaynağı", 28, None, None),
               ("Hesap / Ölçüm Esası", 34, None, None), ("Durum", 30, None, None)])

    ws = wb.create_sheet("Pursantaj Keşif Analizi", index=wb.sheetnames.index("Açık Sorular"))
    baslik(ws, "PURSANTAJ KEŞİF ANALİZİ · ANAHTAR TESLİM SAĞLAMASI",
           "Kurumun pursantaj oranları ile bizim metrajımızdaki grup payları karşılaştırılır. Oranlar yüzde sayısıdır "
           "(3,2164 = %3,2164), Fark yüzde puanıdır. ✔ / ⚠ / ✗ eşikleri Özet sayfasından değiştirilir. Bizim grupların "
           "hangi kurum grubuna sayılacağı Grup Eşleme sayfasında seçilir.", 11)
    GE = "'Grup Eşleme'"
    A, B, C = f"$A${BAS}:$A${SON}", f"$B${BAS}:$B${SON}", f"$C${BAS}:$C${SON}"
    ust = lambda x: f"IFERROR(VLOOKUP({x},$A${BAS}:$B${SON},2,FALSE),\"\")"          # noqa: E731
    kat = lambda x: f"IFERROR(VLOOKUP({x},$A${BAS}:$C${SON},3,FALSE)/100,1)"      # noqa: E731
    p1 = "B{r}"
    p2 = ust(p1)
    p3 = ust(p2)
    tablo(ws, [
        ("İş Grubu (kurum)", 30, None, None),
        ("Üst Grup", 24, None, None),
        ("Oran (üst gruba göre, %)", 12, "0.0000", None),
        ("Genel Oran (%)", 12, "0.0000",
         '=IF(OR(A{r}="",C{r}=""),"",C{r}*' + kat(p1) + "*" + kat(p2) + "*" + kat(p3) + ")"),
        ("Bizim Tutar (₺)", 17, PARA, f'=IF(A{{r}}="","",SUMIF({GE}!$B${BAS}:$B${SON},A{{r}},{GE}!$C${BAS}:$C${SON}))'),
        ("Bizim Pay (%)", 11, "0.0000",
         f'=IF(OR(A{{r}}="",{M}!$K$3=0,COUNTIF({B},A{{r}})>0),"",E{{r}}/{M}!$K$3*100)'),
        ("Fark (puan)", 11, "+0.0000;-0.0000;0", '=IF(OR(D{r}="",F{r}=""),"",F{r}-D{r})'),
        ("Fark (%)", 10, "+0.0;-0.0;0", '=IF(OR(G{r}="",D{r}=0),"",G{r}/D{r}*100)'),
        ("Durum", 22, None,
         f'=IF(A{{r}}="","",IF(COUNTIF({B},A{{r}})>0,IF(ABS(SUMIF({B},A{{r}},{C})-100)<0.0001,'
         f'"✔ alt gruplar Σ=100","⚠ alt gruplar Σ≠100"),IF(G{{r}}="","—",IF(ABS(G{{r}})<=PursUyum,"✔ uyumlu",'
         f'IF(ABS(G{{r}})<=PursCiddi,"⚠ dikkat","✗ ciddi fark")))))'),
        ("Kaynak", 24, None, None),
        ("Not", 34, None, None),
    ])
    # toplam yalnızca alt grubu olmayan (yaprak) gruplardan: üst gruplar çift sayılmasın
    for j, bicim in ((5, PARA), (6, "0.0000")):
        h = get_column_letter(j)
        c = ws.cell(3, j, f'=SUMPRODUCT((COUNTIF({B},{A})=0)*({A}<>""),{h}{BAS}:{h}{SON})')
        c.font = Font(bold=True)
        c.number_format = bicim
    renk(ws, f"I{BAS}:I{SON}", [("✗", KIRMIZI), ("⚠", SARI), ("✔", YESIL)])

    ws = wb.create_sheet("Grup Eşleme", index=wb.sheetnames.index("Açık Sorular"))
    baslik(ws, "GRUP EŞLEME · BİZİM GRUPLAR → KURUM GRUPLARI",
           "Metrajdaki her iş grubu ya da alt başlığın pursantajda hangi kurum grubuna sayılacağı. Karşılığı yoksa "
           "kurum grubu boş bırakılır. Gruplama senaryosu denemek için yalnızca Kurum Grubu değiştirilir.", 8)
    eslesir = f'(({ML("A")}=A{{r}})+({ML("B")}=A{{r}})>0)'
    tablo(ws, [
        ("Bizim Grup (iş grubu / alt başlık)", 30, None, None),
        ("Kurum Grubu", 26, None, None),
        ("Bizim Tutar (₺)", 17, PARA, f'=IF(A{{r}}="","",SUMPRODUCT({eslesir}*{ML("K")}))'),
        ("Pay (%)", 10, "0.0000", f'=IF(OR(A{{r}}="",{M}!$K$3=0),"",C{{r}}/{M}!$K$3*100)'),
        ("Satır", 8, "0", f'=IF(A{{r}}="","",SUMPRODUCT({eslesir}*({ML("C")}<>"")))'),
        ("Miktarı Girilmemiş", 10, "0", f'=IF(A{{r}}="","",SUMPRODUCT({eslesir}*({ML("C")}<>"")*(LEFT({ML("C")},1)<>"—")*({ML("F")}="")))'),
        ("Fiyatı Olmayan", 10, "0", f'=IF(A{{r}}="","",SUMPRODUCT({eslesir}*({ML("F")}<>"")*(LEFT({ML("C")},1)<>"—")*({ML("J")}=0)))'),
        ("Not", 40, None, None),
    ])
    dv = DataValidation(type="list", formula1=f"{P}!$A${BAS}:$A${SON}", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"B{BAS}:B{SON}")

    ws = wb.create_sheet("Kazı Derinlik Analizi")
    baslik(ws, "KAZI DERİNLİK ANALİZİ",
           "Boykesitlerden koşu koşu kazı derinliği: H = (zemin kotu − akar kot) + kazı ek derinliği. İksa, H iksa "
           "eşiğini geçen koşularda iki yüz için 2×H×L; kazı hendek genişliği × H × L. Parametreler Özet sayfasındadır.",
           12)
    tablo(ws, [
        ("Pafta", 8, None, None),
        ("Baca (baş)", 11, None, None),
        ("Baca (son)", 11, None, None),
        ("Zemin−Akar Baş (m)", 11, "0.00", None),
        ("Zemin−Akar Son (m)", 11, "0.00", None),
        ("H1 (m)", 9, "0.00", '=IF(D{r}="","",D{r}+KaziEk)'),
        ("H2 (m)", 9, "0.00", '=IF(E{r}="","",E{r}+KaziEk)'),
        ("H Ort. (m)", 9, "0.000", '=IF(OR(F{r}="",G{r}=""),"",(F{r}+G{r})/2)'),
        ("Uzunluk (m)", 11, "#,##0.00", None),
        ("İksa 2×H×L (m²)", 14, "#,##0.00", '=IF(OR(H{r}="",I{r}=""),"",IF(H{r}>=IksaH,2*H{r}*I{r},0))'),
        ("Kazı B×H×L (m³)", 14, "#,##0.00", '=IF(OR(H{r}="",I{r}=""),"",HendekB*H{r}*I{r})'),
        ("Not", 30, None, None),
    ], toplamlar=(9, 10, 11))

    ortak_sayfalar(wb)
    harita(wb, ORTAK_HARITA + EK_HARITA + [
        ("BFTC", "csv:bftc.csv", "Sıra No", "#sira"),
        ("BFTC", "csv:bftc.csv", "Poz No", "poz_no"),
        ("BFTC", "csv:bftc.csv", "İmalatın Cinsi", "tanim"),
        ("BFTC", "csv:bftc.csv", "Birim", "birim"),
        ("BFTC", "csv:bftc.csv", "Birim Fiyat", "birim_fiyat"),
        ("BFTC", "csv:bftc.csv", "Not", "not"),
        ("Dönemsel BFTC", "csv:donemsel-fiyat.csv", "Sıra No", "#sira"),
        ("Dönemsel BFTC", "csv:donemsel-fiyat.csv", "Poz No", "poz_no"),
        ("Dönemsel BFTC", "csv:donemsel-fiyat.csv", "İmalatın Cinsi", "tanim"),
        ("Dönemsel BFTC", "csv:donemsel-fiyat.csv", "Birim", "birim"),
        *[("Dönemsel BFTC", "csv:donemsel-fiyat.csv", ay, ay.lower()) for ay in AYLAR],
        ("Dönemsel BFTC", "csv:donemsel-fiyat.csv", "Yıllık Liste", "yillik_liste"),
        ("Dönemsel BFTC", "csv:donemsel-fiyat.csv", "Yıllık Kitap", "yillik_kitap"),
        ("Dönemsel BFTC", "csv:donemsel-fiyat.csv", "Kullanılan Dönem", "kullanilan_donem"),
        ("Dönemsel BFTC", "csv:donemsel-fiyat.csv", "Not / Dönem Kaynağı", "not"),
        ("Teknik Tarifler", "csv:teknik-tarifler.csv", "Sıra No", "#sira"),
        ("Teknik Tarifler", "csv:teknik-tarifler.csv", "Poz No", "poz_no"),
        ("Teknik Tarifler", "csv:teknik-tarifler.csv", "Teknik Tarifi", "tarif"),
        ("Metraj Mahal Listesi", "csv:metraj.csv", "İş Grubu", "is_grubu"),
        ("Metraj Mahal Listesi", "csv:metraj.csv", "Alt Başlık", "alt_baslik"),
        ("Metraj Mahal Listesi", "csv:metraj.csv", "Poz No", "poz_no"),
        ("Metraj Mahal Listesi", "csv:metraj.csv", "İmalatın Cinsi", "tanim"),
        ("Metraj Mahal Listesi", "csv:metraj.csv", "Birim", "birim"),
        ("Metraj Mahal Listesi", "csv:metraj.csv", "Miktar", "miktar"),
        ("Metraj Mahal Listesi", "csv:metraj.csv", "Kaynak (çizim / katman / etiket)", "kaynak"),
        ("Metraj Mahal Listesi", "csv:metraj.csv", "Hesap (ölçü kuralı)", "hesap"),
        ("Metraj Mahal Listesi", "csv:metraj.csv", "Durum", "durum"),
        ("Mevcut Mahal Listesi", "csv:mevcut-mahal.csv", "Sıra No", "#sira"),
        ("Mevcut Mahal Listesi", "csv:mevcut-mahal.csv", "Grup", "grup"),
        ("Mevcut Mahal Listesi", "csv:mevcut-mahal.csv", "İmalat / Mevcut Tesis", "imalat"),
        ("Mevcut Mahal Listesi", "csv:mevcut-mahal.csv", "Birim", "birim"),
        ("Mevcut Mahal Listesi", "csv:mevcut-mahal.csv", "Miktar", "miktar"),
        ("Mevcut Mahal Listesi", "csv:mevcut-mahal.csv", "Ölçüm Kaynağı", "olcum_kaynagi"),
        ("Mevcut Mahal Listesi", "csv:mevcut-mahal.csv", "Hesap / Ölçüm Esası", "hesap"),
        ("Mevcut Mahal Listesi", "csv:mevcut-mahal.csv", "Durum", "durum"),
        ("Pursantaj Keşif Analizi", "csv:pursantaj.csv", "İş Grubu (kurum)", "is_grubu"),
        ("Pursantaj Keşif Analizi", "csv:pursantaj.csv", "Üst Grup", "ust_grup"),
        ("Pursantaj Keşif Analizi", "csv:pursantaj.csv", "Oran (üst gruba göre, %)", "oran"),
        ("Pursantaj Keşif Analizi", "csv:pursantaj.csv", "Kaynak", "kaynak"),
        ("Pursantaj Keşif Analizi", "csv:pursantaj.csv", "Not", "not"),
        ("Grup Eşleme", "csv:grup-esleme.csv", "Bizim Grup (iş grubu / alt başlık)", "bizim_grup"),
        ("Grup Eşleme", "csv:grup-esleme.csv", "Kurum Grubu", "kurum_grubu"),
        ("Grup Eşleme", "csv:grup-esleme.csv", "Not", "not"),
        ("Kazı Derinlik Analizi", "csv:kazi-derinlik.csv", "Pafta", "pafta"),
        ("Kazı Derinlik Analizi", "csv:kazi-derinlik.csv", "Baca (baş)", "baca_bas"),
        ("Kazı Derinlik Analizi", "csv:kazi-derinlik.csv", "Baca (son)", "baca_son"),
        ("Kazı Derinlik Analizi", "csv:kazi-derinlik.csv", "Zemin−Akar Baş (m)", "zemin_akar_bas"),
        ("Kazı Derinlik Analizi", "csv:kazi-derinlik.csv", "Zemin−Akar Son (m)", "zemin_akar_son"),
        ("Kazı Derinlik Analizi", "csv:kazi-derinlik.csv", "Uzunluk (m)", "uzunluk"),
        ("Kazı Derinlik Analizi", "csv:kazi-derinlik.csv", "Not", "not"),
    ], "anahtar-teslim")
    wb.save(cikti / "anahtar-teslim.xlsx")


def main() -> None:
    cikti = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "templates" / "excel"
    cikti.mkdir(parents=True, exist_ok=True)
    birim_fiyat(cikti)
    anahtar_teslim(cikti)
    print(f"Şablonlar: {cikti}")


if __name__ == "__main__":
    main()
